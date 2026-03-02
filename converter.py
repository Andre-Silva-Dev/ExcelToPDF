"""
ExcelToPDF - Conversor de Excel para PDF
=========================================
Ferramenta simples para converter arquivos .xlsx e .xls em PDF.
Funciona com Microsoft Excel instalado OU com LibreOffice (gratuito).

Fluxo do programa:
  1. Uma janela abre para você escolher o arquivo Excel.
  2. O programa converte automaticamente para PDF.
  3. A pasta com o PDF é aberta automaticamente ao final.
"""

import os
import platform
import subprocess
import shutil

# Constantes para orientação de página do Microsoft Excel
_XL_PORTRAIT = 1
_XL_LANDSCAPE = 2


# ---------------------------------------------------------------------------
# DETECÇÃO DE ORIENTAÇÃO
# ---------------------------------------------------------------------------

def _orientacao_planilha(caminho_excel: str) -> str:
    """
    Detecta a orientação recomendada para o PDF com base nas dimensões
    da planilha ativa. Retorna 'paisagem' se houver mais colunas do que
    linhas (planilha mais larga), ou 'retrato' caso contrário.

    Funciona apenas com arquivos .xlsx (via openpyxl).
    Para .xls ou em caso de erro, retorna 'paisagem' como padrão conservador
    para evitar que conteúdos largos sejam cortados.
    """
    if os.path.splitext(caminho_excel)[1].lower() != ".xlsx":
        return "paisagem"
    try:
        import openpyxl  # type: ignore[import]
        wb = openpyxl.load_workbook(caminho_excel, read_only=True, data_only=True)
        ws = wb.active
        colunas = ws.max_column or 1
        linhas = ws.max_row or 1
        wb.close()
        return "paisagem" if colunas > linhas else "retrato"
    except Exception:
        return "paisagem"


# ---------------------------------------------------------------------------
# CONVERSÃO: cenário 1 – Microsoft Excel instalado (Windows)
# ---------------------------------------------------------------------------

def converter_com_excel(caminho_excel: str, caminho_pdf: str, orientacao: str = "auto") -> None:
    """
    Converte o arquivo Excel para PDF usando o Microsoft Excel (Windows).

    orientacao: 'auto' (detecta automaticamente), 'retrato' ou 'paisagem'.
    Levanta RuntimeError se algo der errado.
    """
    try:
        import win32com.client  # type: ignore[import]
    except ImportError:
        raise RuntimeError(
            "O módulo 'pywin32' não está instalado.\n"
            "Execute: pip install pywin32"
        )

    excel = None
    pasta = None
    try:
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False

        pasta = excel.Workbooks.Open(caminho_excel)

        # Determina orientação: em modo "auto", decide pela primeira planilha
        if orientacao == "auto":
            primeira = pasta.Sheets(1)
            usado = primeira.UsedRange
            xl_orientacao = (
                _XL_LANDSCAPE if usado.Columns.Count > usado.Rows.Count else _XL_PORTRAIT
            )
        else:
            xl_orientacao = _XL_LANDSCAPE if orientacao == "paisagem" else _XL_PORTRAIT

        for i in range(1, pasta.Sheets.Count + 1):
            planilha = pasta.Sheets(i)
            planilha.PageSetup.Orientation = xl_orientacao
            # Restringe a impressão ao conteúdo preenchido da planilha
            planilha.PageSetup.PrintArea = planilha.UsedRange.Address

        # xlTypePDF = 0
        pasta.ExportAsFixedFormat(0, caminho_pdf)
    except Exception as erro:
        raise RuntimeError(f"Erro ao converter com o Excel: {erro}") from erro
    finally:
        if pasta:
            try:
                pasta.Close(False)
            except Exception:
                pass
        if excel:
            try:
                excel.Quit()
            except Exception:
                pass


# ---------------------------------------------------------------------------
# CONVERSÃO: cenário 2 – LibreOffice instalado (Windows, macOS, Linux)
# ---------------------------------------------------------------------------

def _caminho_libreoffice() -> str | None:
    """
    Localiza o executável do LibreOffice no sistema.
    Retorna o caminho se encontrado, caso contrário None.
    """
    # Verificação pelo PATH (funciona em Linux/macOS)
    if shutil.which("libreoffice"):
        return "libreoffice"
    if shutil.which("soffice"):
        return "soffice"

    # Caminhos comuns no Windows
    caminhos_windows = [
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ]
    for caminho in caminhos_windows:
        if os.path.isfile(caminho):
            return caminho

    # Caminhos comuns no macOS
    caminhos_mac = [
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    ]
    for caminho in caminhos_mac:
        if os.path.isfile(caminho):
            return caminho

    return None


def converter_com_libreoffice(caminho_excel: str, caminho_pdf: str, orientacao: str = "auto") -> None:
    """
    Converte o arquivo Excel para PDF usando o LibreOffice.

    orientacao: 'auto' (detecta automaticamente), 'retrato' ou 'paisagem'.
    Levanta RuntimeError se algo der errado.
    """
    soffice = _caminho_libreoffice()
    if soffice is None:
        raise RuntimeError(
            "LibreOffice não encontrado.\n"
            "Instale o LibreOffice gratuitamente em: https://www.libreoffice.org/download/libreoffice/"
        )

    pasta_destino = os.path.dirname(caminho_pdf)

    # Resolve orientação automática antes de criar arquivo temporário
    if orientacao == "auto":
        orientacao = _orientacao_planilha(caminho_excel)

    # Para .xlsx, aplica a orientação em uma cópia temporária usando openpyxl
    arquivo_a_converter = caminho_excel
    arquivo_temp = None
    if os.path.splitext(caminho_excel)[1].lower() == ".xlsx":
        try:
            import openpyxl  # type: ignore[import]
            import tempfile

            fd, arquivo_temp = tempfile.mkstemp(suffix=".xlsx")
            os.close(fd)
            shutil.copy2(caminho_excel, arquivo_temp)

            wb = openpyxl.load_workbook(arquivo_temp)
            orientacao_openpyxl = (
                "landscape" if orientacao == "paisagem" else "portrait"
            )
            from openpyxl.utils import get_column_letter  # type: ignore[import]
            for ws in wb.worksheets:
                ws.page_setup.orientation = orientacao_openpyxl
                # Restringe a área de impressão ao conteúdo preenchido
                if ws.max_row and ws.max_column:
                    min_col = get_column_letter(ws.min_column or 1)
                    max_col = get_column_letter(ws.max_column)
                    min_row = ws.min_row or 1
                    ws.print_area = f"{min_col}{min_row}:{max_col}{ws.max_row}"
            wb.save(arquivo_temp)
            wb.close()

            arquivo_a_converter = arquivo_temp
        except ImportError:
            pass  # openpyxl não disponível; converte sem ajustar orientação
        except Exception:
            pass  # arquivo inválido ou erro ao aplicar orientação; converte sem ajustar

    try:
        resultado = subprocess.run(
            [
                soffice,
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                pasta_destino,
                arquivo_a_converter,
            ],
            capture_output=True,
            text=True,
            timeout=120,
        )
    except subprocess.TimeoutExpired:
        raise RuntimeError("A conversão demorou demais e foi cancelada.") from None
    except Exception as erro:
        raise RuntimeError(f"Erro ao executar o LibreOffice: {erro}") from erro
    finally:
        if arquivo_temp and os.path.exists(arquivo_temp):
            try:
                os.remove(arquivo_temp)
            except Exception:
                pass

    if resultado.returncode != 0:
        detalhe = resultado.stderr.strip() or resultado.stdout.strip()
        raise RuntimeError(f"O LibreOffice retornou um erro:\n{detalhe}")

    # O LibreOffice cria o PDF com o mesmo nome do arquivo de entrada, extensão .pdf
    nome_base = os.path.splitext(os.path.basename(arquivo_a_converter))[0]
    pdf_gerado = os.path.join(pasta_destino, nome_base + ".pdf")

    # Se o usuário pediu um nome diferente, renomeia
    if os.path.abspath(pdf_gerado) != os.path.abspath(caminho_pdf):
        if os.path.exists(pdf_gerado):
            os.replace(pdf_gerado, caminho_pdf)

    # Confirma que o PDF foi realmente criado (LibreOffice pode retornar 0 sem gerar saída)
    if not os.path.exists(caminho_pdf):
        raise RuntimeError(
            f"LibreOffice indicou sucesso, mas o PDF não foi encontrado em: {caminho_pdf}"
        )


# ---------------------------------------------------------------------------
# DETECÇÃO AUTOMÁTICA DE MÉTODO DE CONVERSÃO
# ---------------------------------------------------------------------------

def _excel_disponivel() -> bool:
    """Verifica se o Microsoft Excel está disponível (somente Windows)."""
    if platform.system() != "Windows":
        return False
    try:
        import win32com.client  # type: ignore[import]  # noqa: F401
        import pythoncom  # type: ignore[import]  # noqa: F401

        pythoncom.CoInitialize()
        try:
            app = win32com.client.Dispatch("Excel.Application")
            app.Quit()
            return True
        finally:
            pythoncom.CoUninitialize()
    except Exception:
        return False


def converter(caminho_excel: str, caminho_pdf: str, orientacao: str = "auto") -> str:
    """
    Escolhe automaticamente o método disponível (Excel ou LibreOffice)
    e realiza a conversão.

    orientacao: 'auto' (detecta automaticamente), 'retrato' ou 'paisagem'.
    Retorna o nome do método utilizado ('Excel' ou 'LibreOffice').
    Levanta RuntimeError em caso de falha.
    """
    if _excel_disponivel():
        converter_com_excel(caminho_excel, caminho_pdf, orientacao)
        return "Excel"

    if _caminho_libreoffice():
        converter_com_libreoffice(caminho_excel, caminho_pdf, orientacao)
        return "LibreOffice"

    raise RuntimeError(
        "Nenhum programa de conversão encontrado.\n\n"
        "Você precisa ter instalado:\n"
        "  • Microsoft Excel  OU\n"
        "  • LibreOffice (gratuito em https://www.libreoffice.org/)\n\n"
        "Após instalar um deles, execute o programa novamente."
    )


# ---------------------------------------------------------------------------
# FEEDBACK VISUAL: abre a pasta do PDF gerado
# ---------------------------------------------------------------------------

def abrir_pasta(caminho_pasta: str) -> None:
    """
    Abre o gerenciador de arquivos na pasta indicada.
    Funciona no Windows, macOS e Linux.
    """
    sistema = platform.system()
    try:
        if sistema == "Windows":
            os.startfile(caminho_pasta)  # type: ignore[attr-defined]
        elif sistema == "Darwin":  # macOS
            subprocess.Popen(["open", caminho_pasta])
        else:  # Linux e outros
            subprocess.Popen(["xdg-open", caminho_pasta])
    except Exception:
        pass  # Falhar ao abrir a pasta não deve interromper o programa


# ---------------------------------------------------------------------------
# VALIDAÇÃO DO ARQUIVO
# ---------------------------------------------------------------------------

def validar_arquivo(caminho: str) -> None:
    """
    Verifica se o arquivo existe e tem extensão suportada.
    Levanta ValueError com mensagem amigável em caso de problema.
    """
    if not caminho:
        raise ValueError("Nenhum arquivo foi selecionado.")

    if not os.path.isfile(caminho):
        raise ValueError(
            f"O arquivo abaixo não foi encontrado:\n{caminho}\n\n"
            "Verifique se o arquivo não foi movido ou renomeado."
        )

    extensao = os.path.splitext(caminho)[1].lower()
    if extensao not in (".xlsx", ".xls"):
        raise ValueError(
            f"O arquivo selecionado não é um Excel válido.\n"
            f"Extensão encontrada: '{extensao}'\n\n"
            "Por favor, selecione um arquivo com extensão .xlsx ou .xls."
        )


# ---------------------------------------------------------------------------
# PONTO DE ENTRADA PRINCIPAL
# ---------------------------------------------------------------------------

def main() -> None:
    """
    Fluxo principal do programa:
      1. Selecionar arquivo Excel via janela de escolha.
      2. Validar o arquivo.
      3. Converter para PDF.
      4. Exibir mensagem de sucesso e abrir a pasta.
    """
    import tkinter as tk
    from tkinter import filedialog, messagebox

    # Inicializa o tkinter uma única vez para as caixas de diálogo
    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)

    # Passo 1: Selecionar arquivo
    messagebox.showinfo(
        "ExcelToPDF",
        "Bem-vindo ao ExcelToPDF!\n\n"
        "Na próxima tela, selecione o arquivo Excel (.xlsx) que deseja converter para PDF.\n\n"
        "Clique em OK para continuar.",
    )

    caminho_excel = filedialog.askopenfilename(
        title="Selecione o arquivo Excel",
        filetypes=[("Arquivos Excel", "*.xlsx *.xls"), ("Todos os arquivos", "*.*")],
    )

    if not caminho_excel:
        messagebox.showwarning("Operação cancelada", "Nenhum arquivo foi selecionado. O programa será encerrado.")
        root.destroy()
        return

    # Passo 2: Validar arquivo
    try:
        validar_arquivo(caminho_excel)
    except ValueError as erro:
        messagebox.showerror("Arquivo inválido", str(erro))
        root.destroy()
        return

    # Passo 3: Definir caminho do PDF e converter
    pasta_destino = os.path.dirname(os.path.abspath(caminho_excel))
    nome_base = os.path.splitext(os.path.basename(caminho_excel))[0]
    caminho_pdf = os.path.join(pasta_destino, nome_base + ".pdf")

    try:
        metodo = converter(caminho_excel, caminho_pdf)
    except RuntimeError as erro:
        messagebox.showerror("Erro na conversão", str(erro))
        root.destroy()
        return

    # Passo 4: Sucesso – abrir pasta e notificar usuário
    abrir_pasta(pasta_destino)

    messagebox.showinfo(
        "Conversão concluída! ✔",
        f"O PDF foi gerado com sucesso usando {metodo}!\n\n"
        f"Arquivo salvo em:\n{caminho_pdf}\n\n"
        "A pasta foi aberta automaticamente para você.",
    )

    root.destroy()


if __name__ == "__main__":
    main()
